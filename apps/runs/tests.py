# pyright: reportAttributeAccessIssue=false, reportMissingImports=false
from datetime import timedelta
from unittest.mock import Mock, patch

from django.test import TestCase, override_settings
from django.utils import timezone
from rest_framework import status
from rest_framework.test import APIClient

from apps.evaluations.models import Evaluation
from apps.llms.models import LLMConfig
from apps.prompts.models import Prompt, PromptPopulation
from apps.responses.models import LLMResponse

from .models import EvaluationRun
from .services import INTERRUPTED_MESSAGE, EvaluationRunService, _live_runs


def _result(text="ok"):
    return {"text": text, "input_tokens": 1, "output_tokens": 1, "latency_ms": 5}


class RunTestMixin:
    def setUp(self):
        self.client = APIClient()
        self.evaluation = Evaluation.objects.create(
            name="Run tests",
            use_case_description="Exercise run execution.",
            task_type=Evaluation.TaskType.TEXT_GENERATION,
        )
        self.population = PromptPopulation.objects.create(
            evaluation=self.evaluation,
            provider="manual",
            model_name="manual",
            requested_prompt_count=3,
            actual_prompt_count=3,
            status=PromptPopulation.Status.COMPLETED,
        )
        self.prompts = [
            Prompt.objects.create(
                population=self.population, text=f"Prompt {sequence}", sequence=sequence
            )
            for sequence in (1, 2, 3)
        ]
        self.llm_config = LLMConfig.objects.create(
            evaluation=self.evaluation,
            provider="openai_compatible",
            model_name="target-model",
            generations_per_prompt=2,
        )

    def make_run(self, **kwargs):
        defaults = {
            "evaluation": self.evaluation,
            "prompt_population": self.population,
            "status": EvaluationRun.Status.RUNNING,
            "total_prompts": 3,
            "total_generations": 6,
        }
        defaults.update(kwargs)
        return EvaluationRun.objects.create(**defaults)


class EvaluationRunServiceTests(RunTestMixin, TestCase):
    def execute(self, run, client):
        with patch("apps.runs.services.build_llm_client", return_value=client):
            EvaluationRunService()._execute(run.id, self.llm_config.id)
        run.refresh_from_db()
        return run

    def test_resume_recomputes_progress_from_distinct_prompts(self):
        run = self.make_run(status=EvaluationRun.Status.FAILED)
        LLMResponse.objects.create(
            run=run, prompt=self.prompts[0], generation_index=1, response_text="kept"
        )
        client = Mock()
        client.generate.return_value = _result()

        run = self.execute(run, client)

        self.assertEqual(run.status, EvaluationRun.Status.COMPLETED)
        self.assertEqual(client.generate.call_count, 5)
        self.assertEqual(run.completed_generations, 6)
        self.assertEqual(run.completed_prompts, 3)
        self.assertEqual(run.responses.count(), 6)
        self.assertIsNotNone(run.heartbeat_at)

    def test_single_generation_failure_is_recorded_without_failing_run(self):
        run = self.make_run()
        client = Mock()
        client.generate.side_effect = [
            _result(), RuntimeError("timeout"), _result(), _result(), _result(), _result()
        ]

        with self.assertLogs("apps.runs.services", level="WARNING"):
            run = self.execute(run, client)

        self.assertEqual(run.status, EvaluationRun.Status.COMPLETED)
        self.assertEqual(run.completed_generations, 5)
        failed = run.responses.get(error__isnull=False)
        self.assertEqual(failed.error, "timeout")
        self.assertEqual(failed.response_text, "")
        status_data = self.client.get(f"/api/runs/{run.id}/status/").data
        self.assertEqual(status_data["failedGenerations"], 1)
        self.assertEqual(
            set(status_data),
            {
                "id", "status", "totalPrompts", "completedPrompts", "totalGenerations",
                "completedGenerations", "failedGenerations", "progress", "errorMessage",
                "external", "heartbeatAt", "startedAt", "finishedAt", "label",
                "llmSnapshot", "promptPopulationId",
            },
        )

    def test_resume_retries_failed_generations(self):
        run = self.make_run(status=EvaluationRun.Status.FAILED)
        LLMResponse.objects.create(
            run=run, prompt=self.prompts[0], generation_index=1, response_text="", error="boom"
        )
        client = Mock()
        client.generate.return_value = _result("retried")

        run = self.execute(run, client)

        self.assertEqual(client.generate.call_count, 6)
        row = run.responses.get(prompt=self.prompts[0], generation_index=1)
        self.assertIsNone(row.error)
        self.assertEqual(row.response_text, "retried")

    @override_settings(TARGET_LLM_FAILFAST_GENERATIONS=3)
    def test_systemic_failure_fails_the_run(self):
        run = self.make_run()
        client = Mock()
        client.generate.side_effect = RuntimeError("connection refused")

        with self.assertLogs("apps.runs.services", level="WARNING"):
            run = self.execute(run, client)

        self.assertEqual(run.status, EvaluationRun.Status.FAILED)
        self.assertIn("connection refused", run.error_message)
        self.assertEqual(client.generate.call_count, 3)
        self.evaluation.refresh_from_db()
        self.assertEqual(self.evaluation.status, Evaluation.Status.FAILED)

    @override_settings(TARGET_LLM_CONCURRENCY=3)
    def test_concurrent_generation_stores_every_result(self):
        run = self.make_run()
        client = Mock()
        client.generate.return_value = _result()

        run = self.execute(run, client)

        self.assertEqual(run.status, EvaluationRun.Status.COMPLETED)
        self.assertEqual(run.completed_generations, 6)
        self.assertEqual(run.responses.count(), 6)

    def test_missing_run_is_logged_not_raised(self):
        with self.assertLogs("apps.runs.services", level="ERROR"):
            EvaluationRunService()._execute(999999, self.llm_config.id)


class RunReapingTests(RunTestMixin, TestCase):
    def test_orphaned_thread_run_is_marked_failed_on_status(self):
        run = self.make_run(heartbeat_at=timezone.now() - timedelta(minutes=5))
        self.evaluation.status = Evaluation.Status.RUNNING
        self.evaluation.save(update_fields=["status"])

        with self.assertLogs("apps.runs.services", level="WARNING"):
            response = self.client.get(f"/api/runs/{run.id}/status/")

        self.assertEqual(response.data["status"], EvaluationRun.Status.FAILED)
        self.assertEqual(response.data["errorMessage"], INTERRUPTED_MESSAGE)
        self.evaluation.refresh_from_db()
        self.assertEqual(self.evaluation.status, Evaluation.Status.FAILED)

    def test_live_or_fresh_runs_are_not_reaped(self):
        live = self.make_run(heartbeat_at=timezone.now() - timedelta(minutes=5))
        _live_runs.add(live.id)
        self.addCleanup(_live_runs.discard, live.id)
        fresh = self.make_run()

        for run in (live, fresh):
            response = self.client.get(f"/api/runs/{run.id}/status/")
            self.assertEqual(response.data["status"], EvaluationRun.Status.RUNNING)

    def test_external_run_is_not_reaped_for_missing_thread(self):
        run = self.make_run(external=True, heartbeat_at=timezone.now() - timedelta(hours=1))

        response = self.client.get(f"/api/runs/{run.id}/status/")

        self.assertEqual(response.data["status"], EvaluationRun.Status.RUNNING)

    def test_stale_external_run_is_eventually_reaped(self):
        run = self.make_run(external=True, heartbeat_at=timezone.now() - timedelta(days=2))

        with self.assertLogs("apps.runs.services", level="WARNING"):
            response = self.client.get(f"/api/runs/{run.id}/status/")

        self.assertEqual(response.data["status"], EvaluationRun.Status.FAILED)


@patch("apps.runs.views.EvaluationRunService.start")
class StartRunViewTests(RunTestMixin, TestCase):
    def url(self):
        return f"/api/evaluations/{self.evaluation.id}/runs/"

    def test_resume_failed_run(self, start):
        run = self.make_run(status=EvaluationRun.Status.FAILED, error_message="boom")

        response = self.client.post(self.url(), {"resumeRunId": run.id}, format="json")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        run.refresh_from_db()
        self.assertEqual(run.status, EvaluationRun.Status.RUNNING)
        self.assertEqual(run.error_message, "")
        start.assert_called_once()

    def test_resume_non_failed_run_conflicts(self, start):
        run = self.make_run(status=EvaluationRun.Status.COMPLETED)

        response = self.client.post(self.url(), {"resumeRunId": run.id}, format="json")

        self.assertEqual(response.status_code, status.HTTP_409_CONFLICT)
        self.assertIn("error", response.data)
        start.assert_not_called()

    def test_start_while_run_in_progress_conflicts(self, start):
        self.make_run()

        response = self.client.post(self.url(), {}, format="json")

        self.assertEqual(response.status_code, status.HTTP_409_CONFLICT)
        self.assertEqual(
            response.data["error"], "A run is already in progress for this evaluation."
        )
        start.assert_not_called()

    def test_start_with_incomplete_population_is_rejected(self, start):
        self.population.status = PromptPopulation.Status.GENERATING
        self.population.save(update_fields=["status"])

        response = self.client.post(
            self.url(), {"promptPopulationId": self.population.id}, format="json"
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        start.assert_not_called()

    def test_start_external_run_marks_it_external(self, start):
        response = self.client.post(self.url(), {"external": True}, format="json")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        run = EvaluationRun.objects.get(pk=response.data["id"])
        self.assertTrue(run.external)
        # 3 prompts × generationsPerPrompt=2 from the evaluation's llm_config.
        self.assertEqual(run.total_generations, 6)
        start.assert_not_called()

    def test_start_records_llm_snapshot_and_label(self, start):
        self.llm_config.api_key = "secret"
        self.llm_config.save(update_fields=["api_key"])

        response = self.client.post(self.url(), {"label": "  gpt run  "}, format="json")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["label"], "gpt run")
        snapshot = response.data["llmSnapshot"]
        self.assertEqual(snapshot["modelName"], "target-model")
        self.assertEqual(snapshot["generationsPerPrompt"], 2)
        self.assertNotIn("secret", str(snapshot))

        # Changing the evaluation's config later does not rewrite history.
        self.llm_config.model_name = "other-model"
        self.llm_config.save(update_fields=["model_name"])
        run = EvaluationRun.objects.get(pk=response.data["id"])
        self.assertEqual(run.llm_snapshot["modelName"], "target-model")

    def test_list_runs_newest_first(self, start):
        older = self.make_run(status=EvaluationRun.Status.COMPLETED)
        newer = self.make_run(status=EvaluationRun.Status.FAILED)
        LLMResponse.objects.create(
            run=newer, prompt=self.prompts[0], generation_index=1, response_text="", error="boom"
        )

        response = self.client.get(self.url())

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual([item["id"] for item in response.data], [newer.id, older.id])
        self.assertEqual(response.data[0]["failedGenerations"], 1)
        self.assertEqual(response.data[1]["failedGenerations"], 0)

    def test_rename_and_delete_run(self, start):
        run = self.make_run(status=EvaluationRun.Status.COMPLETED)
        LLMResponse.objects.create(run=run, prompt=self.prompts[0], generation_index=1, response_text="ok")

        renamed = self.client.patch(f"/api/runs/{run.id}/", {"label": "baseline"}, format="json")
        self.assertEqual(renamed.status_code, status.HTTP_200_OK)
        self.assertEqual(renamed.data["label"], "baseline")

        deleted = self.client.delete(f"/api/runs/{run.id}/")
        self.assertEqual(deleted.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(EvaluationRun.objects.filter(pk=run.id).exists())
        self.assertFalse(LLMResponse.objects.filter(run_id=run.id).exists())

    def test_active_run_cannot_be_deleted(self, start):
        run = self.make_run()

        response = self.client.delete(f"/api/runs/{run.id}/")

        self.assertEqual(response.status_code, status.HTTP_409_CONFLICT)
        self.assertTrue(EvaluationRun.objects.filter(pk=run.id).exists())


class RunExportTests(RunTestMixin, TestCase):
    def setUp(self):
        super().setUp()
        from apps.fairness.models import LangFairMetricResult

        self.run = self.make_run(status=EvaluationRun.Status.COMPLETED, label="中文標籤")
        self.prompts[0].metadata = {
            "persona_id": "p1",
            "persona_attributes": {"gender_identity": "Woman", "region": "Asia"},
            "fairness": {"pair_id": "cf_1", "role": "original", "group": "female"},
        }
        self.prompts[0].save()
        LLMResponse.objects.create(
            run=self.run, prompt=self.prompts[0], generation_index=1, response_text="回答 one"
        )
        LLMResponse.objects.create(
            run=self.run, prompt=self.prompts[1], generation_index=1, response_text="", error="boom"
        )
        LangFairMetricResult.objects.create(
            run=self.run,
            cache_key="k",
            status=LangFairMetricResult.Status.COMPLETED,
            result={"suites": {"toxicity": {
                "metrics": {"Toxic Fraction": 0.5},
                "confidenceIntervals": {"Toxic Fraction": {"ciLow": 0.1, "ciHigh": 0.9, "n": 2}},
                "groupComparison": {"gender": {
                    "comparison": "female - male",
                    "counts": {"female": {"responses": 1, "units": 1}},
                    "byGroup": {"female": {"Toxic Fraction": 0.5}},
                    "gap": {"metrics": {"Toxic Fraction": 0.2},
                            "confidenceIntervals": {"Toxic Fraction": {"ciLow": -0.1, "ciHigh": 0.4}}},
                }},
            }}},
        )

    def read_csv(self, table):
        import csv
        import io

        response = self.client.get(f"/api/runs/{self.run.id}/export/", {"table": table})
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertTrue(response.content.startswith(b"\xef\xbb\xbf"))  # BOM for Excel
        return list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))

    def test_workbook_has_one_sheet_per_table(self):
        import io

        from openpyxl import load_workbook

        from .export import TABLES

        response = self.client.get(f"/api/runs/{self.run.id}/export/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertTrue(response["Content-Disposition"].endswith('.xlsx"'))
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, [title for _, title, _ in TABLES.values()])
        responses = list(workbook["回應"].iter_rows(values_only=True))
        header = responses[0]
        self.assertEqual(responses[1][header.index("response_text")], "回答 one")

    def test_tables(self):
        responses = self.read_csv("responses")
        self.assertEqual([row["response_text"] for row in responses], ["回答 one", ""])
        self.assertEqual(responses[0]["pair_id"], "cf_1")
        self.assertEqual(responses[0]["gender_group"], "female")
        self.assertEqual(responses[1]["error"], "boom")

        personas = self.read_csv("personas")
        self.assertEqual(personas[0]["persona_id"], "p1")
        self.assertEqual(personas[0]["region"], "Asia")

        metrics = self.read_csv("metrics")
        self.assertEqual(metrics[0]["metric"], "Toxic Fraction")
        self.assertEqual((metrics[0]["ci_low"], metrics[0]["ci_high"]), ("0.1", "0.9"))

        groups = self.read_csv("group_comparison")
        self.assertEqual([row["type"] for row in groups], ["group", "gap"])
        self.assertEqual(groups[1]["ci_low"], "-0.1")

        run_info = {(row["section"], row["field"]): row["value"] for row in self.read_csv("run")}
        self.assertEqual(run_info[("run", "label")], "中文標籤")
        self.assertEqual(run_info[("run", "failed_generations")], "1")

    def test_unknown_table_is_rejected(self):
        response = self.client.get(f"/api/runs/{self.run.id}/export/", {"table": "nope"})
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
